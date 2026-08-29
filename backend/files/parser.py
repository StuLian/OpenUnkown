"""上传文件的解析：把常见文档/表格转成纯文本，供对话附件使用。

支持的格式：
- 纯文本 / Markdown：直接解码；
- CSV：用 csv.Sniffer 探测分隔符，转成易读的表格文本；
- PDF：pypdf 逐页提取文本；
- Word (.docx)：python-docx 提取段落与表格；
- Excel (.xlsx)：openpyxl 按工作表逐行转文本。

解析结果为纯文本，统一做长度截断，避免超大文件把模型上下文撑爆。
pypdf / python-docx / openpyxl 均为可选依赖，按需在函数内延迟 import，
缺少对应库时抛出 ParseError，不影响应用启动。
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

# 支持的文件扩展名（不含点）
ALLOWED_EXTENSIONS = ("txt", "md", "markdown", "csv", "pdf", "docx", "xlsx")

# 单文件大小上限（字节）
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
# 解析后注入上下文的文本长度上限（字符）
MAX_CONTENT_CHARS = 50_000
# CSV / XLSX 最大处理行数，避免超大表格拖慢解析
MAX_TABLE_ROWS = 2_000
# PDF 最多提取的页数
MAX_PDF_PAGES = 200


class ParseError(Exception):
    """文件解析失败（格式不支持 / 内容损坏 / 超出限制）。"""


@dataclass
class ParsedFile:
    """单文件解析结果。"""

    filename: str
    file_type: str  # txt | markdown | csv | pdf | docx | xlsx
    size: int       # 原始字节数
    content: str    # 解析并截断后的纯文本
    char_count: int # 截断前的字符数
    truncated: bool


def _fmt_size(n: int) -> str:
    """把字节数格式化为易读字符串。"""
    if n < 1024:
        return f"{n} B"
    if n < 1024 * 1024:
        return f"{n / 1024:.1f} KB"
    return f"{n / (1024 * 1024):.1f} MB"


def detect_type(filename: str) -> str:
    """按扩展名归一化文件类型；不支持的扩展名抛出 ParseError。"""
    ext = Path(filename or "").suffix.lstrip(".").lower()
    if ext in ("md", "markdown"):
        return "markdown"
    if ext in ALLOWED_EXTENSIONS:
        return ext
    allowed = ", ".join(ALLOWED_EXTENSIONS)
    raise ParseError(f"不支持的文件类型「.{ext or '?'}」，仅支持：{allowed}")


def _decode_text(data: bytes) -> str:
    """按常见编码顺序解码文本，保证中文/英文/GBK 等尽量可读。"""
    for enc in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_txt(data: bytes) -> str:
    return _decode_text(data)


def _parse_csv(data: bytes) -> str:
    text = _decode_text(data)
    try:
        sample = text[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error:
        raise ParseError("CSV 内容无法解析，请确认文件编码与格式") from None
    if not rows:
        return "(空 CSV 文件)"

    header = [c.strip() for c in rows[0]]
    width = max(len(header), *(len(r) for r in rows[:MAX_TABLE_ROWS]))
    lines = [f"CSV 表格（共 {len(rows)} 行）："]
    lines.append(" | ".join(header + [""] * (width - len(header))))
    lines.append("-" * 40)
    for row in rows[1:MAX_TABLE_ROWS]:
        cells = [c for c in row] + [""] * (width - len(row))
        lines.append(" | ".join(cells[:width]))
    if len(rows) > MAX_TABLE_ROWS:
        lines.append(f"...（其余 {len(rows) - MAX_TABLE_ROWS} 行已省略）")
    return "\n".join(lines)


def _parse_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - 依赖缺失时提示
        raise ParseError("缺少 PDF 解析依赖 pypdf，请安装后重试") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                pass
        pages: list[str] = []
        for i, page in enumerate(reader.pages):
            if i >= MAX_PDF_PAGES:
                pages.append(f"...（其余页已省略，共 {len(reader.pages)} 页）")
                break
            pages.append(page.extract_text() or "")
        return "\n\n".join(pages)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PDF 解析失败（文件可能已损坏或加密）: {exc}") from exc


def _parse_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少 Word 解析依赖 python-docx，请安装后重试") from exc
    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Word 解析失败（文件可能已损坏）: {exc}") from exc


def _parse_xlsx(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少 Excel 解析依赖 openpyxl，请安装后重试") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        for ws in wb.worksheets:
            parts.append(f"[工作表: {ws.title}]")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_TABLE_ROWS:
                    parts.append("...(更多行已省略)")
                    break
                parts.append(" | ".join("" if c is None else str(c) for c in row))
            parts.append("")
        wb.close()
        return "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Excel 解析失败（文件可能已损坏）: {exc}") from exc


_PARSERS = {
    "txt": _parse_txt,
    "markdown": _parse_txt,
    "csv": _parse_csv,
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
}


def parse_bytes(filename: str, data: bytes) -> ParsedFile:
    """解析上传文件的字节内容，返回截断后的纯文本与元信息。"""
    if not data:
        raise ParseError("文件内容为空")
    if len(data) > MAX_FILE_SIZE:
        raise ParseError(
            f"文件过大（{_fmt_size(len(data))}），上限 {_fmt_size(MAX_FILE_SIZE)}"
        )
    file_type = detect_type(filename)
    raw = _PARSERS[file_type](data).strip()
    if not raw:
        raw = "(未解析到文本内容)"

    char_count = len(raw)
    truncated = char_count > MAX_CONTENT_CHARS
    content = raw[:MAX_CONTENT_CHARS]
    if truncated:
        content += f"\n\n[内容过长已截断，仅保留前 {MAX_CONTENT_CHARS} 字符]"

    return ParsedFile(
        filename=Path(filename or "").name or "未命名文件",
        file_type=file_type,
        size=len(data),
        content=content,
        char_count=char_count,
        truncated=truncated,
    )
